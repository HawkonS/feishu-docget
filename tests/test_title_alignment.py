import os
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.converters.docx.cleaner import clean_document


class TitleAlignmentTests(unittest.TestCase):
    def _make_document(self, path, alignment):
        document = Document()
        paragraph = document.add_paragraph('标题', style='Heading 1')
        paragraph.alignment = alignment
        document.add_paragraph('正文')
        document.save(path)

    def test_forced_title_alignment_overrides_template_paragraph_alignment(self):
        with tempfile.TemporaryDirectory() as workspace:
            template_path = os.path.join(workspace, 'template.docx')
            output_path = os.path.join(workspace, 'output.docx')
            self._make_document(template_path, WD_ALIGN_PARAGRAPH.RIGHT)
            self._make_document(output_path, WD_ALIGN_PARAGRAPH.RIGHT)

            clean_document(output_path, template_path=template_path, title_align='left')

            document = Document(output_path)
            self.assertEqual(document.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_justify_alignment_is_supported(self):
        with tempfile.TemporaryDirectory() as workspace:
            output_path = os.path.join(workspace, 'output.docx')
            self._make_document(output_path, WD_ALIGN_PARAGRAPH.LEFT)

            clean_document(output_path, title_align='justify')

            document = Document(output_path)
            self.assertEqual(document.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)

    def test_none_keeps_existing_title_alignment(self):
        with tempfile.TemporaryDirectory() as workspace:
            output_path = os.path.join(workspace, 'output.docx')
            self._make_document(output_path, WD_ALIGN_PARAGRAPH.RIGHT)

            clean_document(output_path, title_align='none')

            document = Document(output_path)
            self.assertEqual(document.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)


if __name__ == '__main__':
    unittest.main()
