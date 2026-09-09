import os
import tempfile
import unittest
import zipfile
from io import BytesIO

from docx import Document
from docx.shared import Cm
from PIL import Image

from src.converters.docx.cleaner import clean_document, _normalize_image_border
from src.converters.docx.converter import _apply_feishu_image_crop
from src.core.image_processor import calculate_center_crop


class ImageStyleTests(unittest.TestCase):
    def _make_document(self, path):
        document = Document()
        run = document.add_paragraph().add_run()
        run.add_picture(os.path.join(os.path.dirname(os.path.dirname(__file__)), 'template', 'Hawkon.png'), width=Cm(5))
        document.save(path)

    def test_image_border_normalizes_frontend_values(self):
        rule = _normalize_image_border({
            'borderEnabled': True,
            'borderColor': '#12abEF',
            'borderWidth': 2,
            'shrinkPercent': 10,
        })
        self.assertEqual(rule, {'color': '12ABEF', 'width_pt': 2.0, 'shrink_percent': 10.0})

    def test_image_border_and_shrink_override_picture_properties(self):
        with tempfile.TemporaryDirectory() as workspace:
            output_path = os.path.join(workspace, 'output.docx')
            self._make_document(output_path)
            clean_document(output_path, image_style={
                'maxWidth': None,
                'maxHeight': None,
                'align': 'center',
                'borderEnabled': True,
                'borderColor': '#112233',
                'borderWidth': 2,
                'shrinkPercent': 10,
            })

            with zipfile.ZipFile(output_path) as archive:
                xml = archive.read('word/document.xml').decode('utf-8')
                padded_images = [name for name in archive.namelist() if name.startswith('word/media/')]
                self.assertGreaterEqual(len(padded_images), 2)
                # The padded replacement retains the original canvas size while
                # its non-transparent content is reduced to 90%.
                padded = Image.open(BytesIO(archive.read(padded_images[-1]))).convert('RGBA')
                alpha = padded.getchannel('A')
                bbox = alpha.getbbox()
                self.assertIsNotNone(bbox)
                self.assertAlmostEqual((bbox[2] - bbox[0]) / padded.width, 0.9, delta=0.03)
            self.assertIn('cx="1800000"', xml)  # outer frame is retained for whitespace
            self.assertIn('w="25400"', xml)  # 2pt in DrawingML EMU
            self.assertIn('srgbClr val="112233"', xml)
            self.assertIn('prstDash val="solid"', xml)

    def test_center_crop_calculation_uses_feishu_frame_ratio(self):
        crop = calculate_center_crop(400, 200, 100, 100)
        self.assertEqual(crop, {
            'left': 0.25,
            'top': 0.0,
            'right': 0.25,
            'bottom': 0.0,
        })
        self.assertIsNone(calculate_center_crop(400, 200, 200, 100))

    def test_feishu_crop_is_written_as_native_docx_source_rectangle(self):
        with tempfile.TemporaryDirectory() as workspace:
            image_path = os.path.join(workspace, 'wide.png')
            output_path = os.path.join(workspace, 'cropped.docx')
            Image.new('RGB', (400, 200), color='red').save(image_path)

            document = Document()
            shape = document.add_paragraph().add_run().add_picture(image_path, width=Cm(10))
            self.assertTrue(_apply_feishu_image_crop(
                shape,
                image_path,
                {'width': 100, 'height': 100},
            ))
            document.save(output_path)

            with zipfile.ZipFile(output_path) as archive:
                xml = archive.read('word/document.xml').decode('utf-8')
            self.assertIn('<a:srcRect l="25000" r="25000"/>', xml)
            self.assertIn('cx="3600000" cy="3600000"', xml)

    def test_uncropped_feishu_image_keeps_original_picture_xml(self):
        with tempfile.TemporaryDirectory() as workspace:
            image_path = os.path.join(workspace, 'wide.png')
            Image.new('RGB', (400, 200), color='blue').save(image_path)

            document = Document()
            shape = document.add_paragraph().add_run().add_picture(image_path, width=Cm(10))
            self.assertFalse(_apply_feishu_image_crop(
                shape,
                image_path,
                {'width': 200, 'height': 100},
            ))
            self.assertNotIn('srcRect', shape._inline.xml)


if __name__ == '__main__':
    unittest.main()
