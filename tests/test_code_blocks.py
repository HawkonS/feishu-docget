import os
import tempfile
import unittest
from unittest.mock import Mock

from docx import Document
from docx.oxml.ns import qn

from src.converters.docx.cleaner import clean_document
from src.converters.docx.converter import (
    FeishuDocxConverter,
    extract_code_block_names_from_content,
)
from src.services.doc_service import _process_document_with_client


CODE = (
    '# Licensed under the Apache License, Version 2.0.\n'
    '"""Hourly governance with all fifteen historical covariates and missing cells."""\n\n'
    'import csv\nimport hashlib\nimport json\nimport os\nfrom pathlib import Path\n'
)


def code_blocks():
    return [
        {'block_id': 'page', 'block_type': 1, 'children': ['code']},
        {
            'block_id': 'code',
            'parent_id': 'page',
            'block_type': 14,
            'code': {'elements': [{'text_run': {'content': CODE}}]},
        },
    ]


class CodeBlockNameTests(unittest.TestCase):
    def test_extracts_rendered_code_caption_by_block_id(self):
        content = (
            '<pre id="code-1" caption="code/governance.py&#xA;" lang="python">'
            '<code>print(1)</code></pre>'
            '<pre id="code-2" caption="" lang="python"><code>print(2)</code></pre>'
            '<pre caption="missing-id"><code>print(3)</code></pre>'
        )
        self.assertEqual(
            extract_code_block_names_from_content(content),
            {'code-1': 'code/governance.py'},
        )

    def test_missing_or_malformed_metadata_does_not_break_export(self):
        for content in (None, '', '<pre caption="broken">'):
            with self.subTest(content=content):
                self.assertEqual(extract_code_block_names_from_content(content), {})

    def test_bold_name_and_code_body_survive_style_cleaning_in_separate_cells(self):
        with tempfile.TemporaryDirectory() as workspace:
            output_path = os.path.join(workspace, 'code.docx')
            converter = FeishuDocxConverter(
                code_blocks(),
                None,
                workspace,
                code_block_names={'code': 'code/governance.py'},
            )
            converter.process(output_path)
            clean_document(output_path, code_block_config={
                'fontFamily': 'Consolas',
                'fontSize': 11,
                'fontColor': '#123456',
                'bgColor': '#EEEEEE',
                'lineSpacing': 1.2,
            })

            document = Document(output_path)
            table = document.tables[0]
            self.assertEqual(len(table.rows), 2)
            self.assertEqual(table.cell(0, 0).text, 'code/governance.py')
            self.assertEqual(table.cell(1, 0).text, CODE)
            self.assertTrue(table.cell(0, 0).paragraphs[0].runs[0].bold)
            self.assertFalse(table.cell(1, 0).paragraphs[0].runs[0].bold)
            for row in table.rows:
                cell = row.cells[0]
                run = cell.paragraphs[0].runs[0]
                self.assertEqual(run.font.name, 'Consolas')
                self.assertEqual(run.font.size.pt, 11)
                self.assertEqual(str(run.font.color.rgb), '123456')
                self.assertEqual(cell._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')), 'EEEEEE')

    def test_unnamed_code_retains_single_cell_and_original_content(self):
        with tempfile.TemporaryDirectory() as workspace:
            output_path = os.path.join(workspace, 'code.docx')
            FeishuDocxConverter(code_blocks(), None, workspace).process(output_path)
            table = Document(output_path).tables[0]
            self.assertEqual(len(table.rows), 1)
            self.assertEqual(table.cell(0, 0).text, CODE)

    def _client(self):
        client = Mock()
        client.extract_doc_id.return_value = 'page'
        client.get_document_meta.return_value = {'title': 'Code'}
        client.get_blocks.return_value = code_blocks()
        return client

    def test_download_fetches_names_for_document_with_only_code(self):
        client = self._client()
        client.get_document_content.return_value = (
            '<pre id="code" caption="code/governance.py&#xA;" lang="python">'
            '<code>ignored: content still comes from the blocks API</code></pre>'
        )
        with tempfile.TemporaryDirectory() as workspace:
            result = _process_document_with_client(client, 'document-url', output_root=workspace)
            client.get_document_content.assert_called_once_with('page')
            table = Document(result['docx_path']).tables[0]
            self.assertEqual(table.cell(0, 0).text, 'code/governance.py')
            self.assertTrue(table.cell(0, 0).paragraphs[0].runs[0].bold)
            self.assertEqual(table.cell(1, 0).text, CODE)

    def test_download_keeps_code_when_caption_endpoint_is_unavailable(self):
        client = self._client()
        client.get_document_content.side_effect = PermissionError('missing scope')
        with tempfile.TemporaryDirectory() as workspace:
            result = _process_document_with_client(client, 'document-url', output_root=workspace)
            table = Document(result['docx_path']).tables[0]
            self.assertEqual(len(table.rows), 1)
            self.assertEqual(table.cell(0, 0).text, CODE)


if __name__ == '__main__':
    unittest.main()
